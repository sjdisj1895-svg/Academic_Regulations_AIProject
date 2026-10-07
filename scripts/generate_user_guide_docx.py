# -*- coding: utf-8 -*-
"""
비개발자(전 부서 공문 배포용) 사용자 가이드 -> .docx 생성 스크립트
실행: python scripts/generate_user_guide_docx.py
(사전에 C:/tmp/manual_home.png, manual_results.png, manual_preview.png 스크린샷 필요)
"""
import os

from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_PATH = os.path.join(ROOT, "경상국립대학교_규정통합검색_사용자안내.docx")
LOGO = os.path.join(ROOT, "web", "assets", "gnu_emblem_b.png")
SCREEN_HOME = "C:/tmp/manual_home.png"
SCREEN_RESULTS = "C:/tmp/manual_results.png"
SCREEN_PREVIEW = "C:/tmp/manual_preview.png"

GNU_BLUE = RGBColor(0x00, 0x9E, 0xDB)
GNU_GREY = RGBColor(0x43, 0x52, 0x5A)


def set_korean_font(run, size=10.5, bold=False, color=None):
    run.font.name = "malgun gothic"
    r = run._element
    r.rPr.rFonts.set(qn("w:eastAsia"), "malgun gothic")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading(doc, text, size=16, color=GNU_BLUE, space_before=18):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_korean_font(run, size=size, bold=True, color=color)
    return p


def add_sub(doc, text, size=12, bold=True, color=GNU_GREY):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_korean_font(run, size=size, bold=bold, color=color)
    return p


def add_body(doc, text, size=10.5, bold=False, bullet=False):
    p = doc.add_paragraph(style="List Bullet" if bullet else None)
    run = p.add_run(text)
    set_korean_font(run, size=size, bold=bold)
    return p


def add_step(doc, no, text, size=10.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(f"{no}. ")
    set_korean_font(run, size=size, bold=True, color=GNU_BLUE)
    run2 = p.add_run(text)
    set_korean_font(run2, size=size)
    return p


def add_image(doc, path, width_cm=15.5, caption=None):
    if not os.path.exists(path):
        add_body(doc, f"[스크린샷 누락: {path}]")
        return
    doc.add_picture(path, width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(caption)
        set_korean_font(run, size=9.5, color=GNU_GREY)
        run.italic = True


def add_table(doc, headers, rows, col_widths_cm=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        set_korean_font(run, size=10, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        tcPr = hdr_cells[i]._tc.get_or_add_tcPr()
        shd = tcPr.makeelement(qn("w:shd"), {
            qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): "009EDB"})
        tcPr.append(shd)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(str(val))
            set_korean_font(run, size=10)
    if col_widths_cm:
        for row in table.rows:
            for i, w in enumerate(col_widths_cm):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return table


def main():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "malgun gothic"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "malgun gothic")
    style.font.size = Pt(10.5)

    section = doc.sections[0]
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    # ---------------------------------------------------------------- 표지
    if os.path.exists(LOGO):
        doc.add_picture(LOGO, width=Cm(3.2))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    run = p.add_run("경상국립대학교 규정 통합 검색 서비스")
    set_korean_font(run, size=22, bold=True, color=GNU_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("이용 안내")
    set_korean_font(run, size=16, bold=True, color=GNU_GREY)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    run = p.add_run("대학·산학협력단 규정을 한 번에 검색할 수 있습니다")
    set_korean_font(run, size=11, color=GNU_GREY)

    doc.add_page_break()

    # ---------------------------------------------------------------- 1. 소개
    add_heading(doc, "1. 서비스 소개")
    add_body(doc, "경상국립대학교의 대학 규정(학칙·규정·지침)과 산학협력단 규정을 "
                  "조(條)·항(項) 단위로 한 번에 검색할 수 있는 통합 검색 서비스입니다.")
    add_body(doc, "그동안 여러 게시판·문서에 흩어져 있던 규정을 일일이 찾아보던 번거로움을 줄이고, "
                  "키워드 하나로 관련 조항을 바로 찾아볼 수 있습니다.", bullet=True)
    add_body(doc, "검색 결과에서 바로 담당부서·연락처를 확인할 수 있어, "
                  "문의할 부서를 찾는 시간도 줄어듭니다.", bullet=True)

    # ---------------------------------------------------------------- 2. 접속
    add_heading(doc, "2. 접속 방법")
    add_table(
        doc,
        ["항목", "내용"],
        [
            ["접속 주소", "https://regulations.gnu.ac.kr"],
            ["로그인", "별도 로그인·회원가입 없이 바로 이용 가능"],
            ["이용 환경", "PC·태블릿·휴대폰 어디서나 접속 가능 (별도 앱 설치 불필요)"],
        ],
        col_widths_cm=[4, 11.5],
    )

    # ---------------------------------------------------------------- 3. 검색
    add_heading(doc, "3. 검색 방법")

    add_sub(doc, "① 검색창에 키워드 입력")
    add_body(doc, "메인 화면 가운데 검색창에 찾고 싶은 규정명이나 단어를 입력하고 "
                  "검색 버튼(돋보기)을 누르거나 Enter를 칩니다.")
    add_image(doc, SCREEN_HOME, caption="그림 1. 메인(검색) 화면")

    add_sub(doc, "② 필터로 범위 좁히기")
    add_step(doc, 1, "'전체 / 대학 / 산학협력단' 중 원하는 출처를 선택합니다.")
    add_step(doc, 2, "그 아래 분류 칩(학칙·규정·지침 또는 제1편~제7편)을 눌러 범위를 더 좁힐 수 있습니다.")
    add_step(doc, 3, "'필터 더보기'에서 폐지된 규정 포함 여부, 최근 개정만 보기 등 세부 옵션을 쓸 수 있습니다.")
    add_image(doc, SCREEN_RESULTS, caption="그림 2. 검색 결과 화면 (예: '휴학' 검색)")

    add_sub(doc, "③ 조항 클릭 → 미리보기 → 원문 확인")
    add_step(doc, 1, "검색 결과에서 원하는 조항을 클릭하면 오른쪽(모바일은 하단)에 미리보기 창이 열립니다.")
    add_step(doc, 2, "미리보기 상단에서 담당부서와 전화번호를 바로 확인할 수 있습니다 (클릭 시 바로 전화 연결).")
    add_step(doc, 3, "'규정 전문' 탭을 누르면 이 조항이 속한 규정 전체 내용을 볼 수 있습니다.")
    add_step(doc, 4, "'관련 사이트로 이동'(또는 '규정집 원문 파일로 이동') 버튼을 누르면 원본 게시글로 이동합니다.")
    add_image(doc, SCREEN_PREVIEW, caption="그림 3. 조항 미리보기 화면 (담당부서·연락처 표시)")

    # ---------------------------------------------------------------- 4. FAQ
    add_heading(doc, "4. 자주 묻는 질문")
    add_sub(doc, "Q. 검색을 했는데 결과가 안 나와요.")
    add_body(doc, "검색어 철자를 확인해보시고, 적용 중인 필터(출처·분류)를 '모두 해제'한 뒤 "
                  "다시 검색해보세요. 너무 긴 문장보다는 핵심 단어 위주로 검색하면 더 잘 찾아집니다.", bullet=True)
    add_sub(doc, "Q. 우리 부서 규정이 안 보여요.")
    add_body(doc, "기본값은 '현재 시행 중인 규정'만 보여줍니다. 폐지되었거나 개정 전 규정을 찾으신다면 "
                  "'필터 더보기 → 폐지된 규정도 포함'을 체크해보세요.", bullet=True)
    add_sub(doc, "Q. 모바일에서도 똑같이 쓸 수 있나요?")
    add_body(doc, "네, 화면이 자동으로 모바일 크기에 맞춰지며 기능은 PC와 동일합니다.", bullet=True)

    # ---------------------------------------------------------------- 5. 문의처
    add_heading(doc, "5. 문의처")
    add_table(
        doc,
        ["담당", "연락처"],
        [
            ["정보전산처", "055-772-0627"],
        ],
        col_widths_cm=[7.5, 7.5],
    )

    doc.save(OUT_PATH)
    print(f"저장 완료: {OUT_PATH}")


if __name__ == "__main__":
    main()
